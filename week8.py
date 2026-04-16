# Week 8 - Excel Output and Advanced Visualization
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

pd.set_option("display.max_columns", None)

# Load data and recreate user profile from Week 7
logs = pd.read_csv("capstone_logs_clean.csv")
logs["timestamp"] = pd.to_datetime(logs["timestamp"])

profiles = pd.DataFrame({
    "user_id": ["USR_001", "USR_002", "USR_003", "USR_004", "USR_005",
                "USR_006", "USR_007", "USR_008", "USR_009", "USR_010"],
    "account_type": ["free", "premium", "free", "premium", "free",
                     "premium", "free", "premium", "free", "premium"],
    "account_age_days": [45, 365, 180, 90, 270, 120, 30, 450, 60, 200],
    "verified": [True, True, False, True, False,
                 True, False, True, False, True]
})

merged = logs.merge(profiles, on="user_id", how="inner")
merged["hour"] = merged["timestamp"].dt.hour
merged["is_nighttime"] = merged["hour"].between(0,4)
merged["is_weekend"] = merged["timestamp"].dt.dayofweek.isin([5,6])

def classify_query(query_text):
    text = query_text.lower()
    if any(word in text for word in ["synthesize", "compound", "stabilize"]):
        return "Chemical"
    elif any(word in text for word in ["pathogen", "transmission"]):
        return "Biological"
    elif any(word in text for word in ["nuclear", "fission"]):
        return "Radiological"
    elif any(word in text for word in ["dispersal", "precursor"]):
        return "Explosive"
    else:
        return "Benign"

merged["cbrn_category"] = merged["query_text"].apply(classify_query)

user_profile = merged.groupby("user_id").agg(
    total_queries=("user_id", "count"),
    avg_response=("response_length", "mean"),
    nighttime_queries=("is_nighttime", "sum"),
    weekend_queries=("is_weekend", "sum"),
    unknown_country=("country", lambda x: (x == "Unknown").sum()),
    cbrn_queries=("cbrn_category", lambda x: (x != "Benign").sum())
).round(1).reset_index()

user_profile["nighttime_pct"] = (
    user_profile["nighttime_queries"] / user_profile["total_queries"] * 100
).round(1)

user_profile["cbrn_pct"] = (
    user_profile["cbrn_queries"] / user_profile["total_queries"] * 100
).round(1)

user_profile = user_profile.merge(
    profiles[["user_id", "account_type", "verified", "account_age_days"]],
    on="user_id"
)

print("Data prepared successfully")
print(f"User profiles: {len(user_profile)}")

# Writing Excel output
print("\n--- Writing Excel Report ---")

with pd.ExcelWriter("investigation_report.xlsx", engine="openpyxl") as writer:

    # Sheet 1 -- User Behavioral Profiles
    user_profile.to_excel(writer, sheet_name="User Profiles",
                          index=False, startrow=1)
    
    ws = writer.sheets["User Profiles"]
    wb = writer.book

    # Add title.row
    ws.cell(row=1, column=1).value = "INVESTIGATION REPORT --- User Behavioral Profiles"
    ws.cell(row=1, column=1).font = openpyxl.styles.Font(
        bold=True, size=14, color="FFFFFF")
    ws.cell(row=1, column=1).fill = openpyxl.styles.PatternFill(
        fill_type="solid", fgColor="1F3864")
    
    # Merge title cells across all columns
    ws.merge_cells("A1:L1")

    # Fix background color
    from openpyxl.styles import PatternFill
    blue_fill = PatternFill(start_color="1F3864",
                            end_color="1F3864",
                            fill_type="solid")
    ws["A1"].fill = blue_fill
    ws["A1"].alignment = openpyxl.styles.Alignment(horizontal="center")
    
    # Set column widths
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 14
    ws.column_dimensions["I"].width = 10
    ws.column_dimensions["J"].width = 12
    ws.column_dimensions["K"].width = 10
    ws.column_dimensions["L"].width = 16

    # Sheet 2 -- Raw logs sample
    logs.head(100).to_excel(writer, sheet_name="Log Sample", index=False)

    # Sheet 3 -- embed the dashboard chart
    ws_chart = wb.create_sheet("Dashboard")
    img = openpyxl.drawing.image.Image("investigation_dashboard.png")
    img.width = 900
    img.height = 600
    ws_chart.add_image(img, "A1")
    print("Dashboard chart embedded in Excel")

    print("Excel report saved as investigation_report.xlsx")

print("Done!")    

# Advanced visualization -- multiple charts
print("\n--- Generating Investigation Charts ---")

fig, axes = plt.subplots(2, 2, figsize=(14, 10))
fig.suptitle("Investigation Dashboard", fontsize=16, fontweight="bold")

# Chart 1 -- CBRN percentage by uder (horizontal bar)
axes[0, 0].barh(user_profile["user_id"],
                user_profile["cbrn_pct"],
                color=["red" if p > 60 else "steelblue"
                       for p in user_profile["cbrn_pct"]])
axes[0, 0].axvline(x=60, color="orange", linestyle="--", label="60% threshold")
axes[0, 0].set_title("CBRN Query Percentage by User")
axes[0, 0].set_xlabel("CBRN %")
axes[0, 0].legend()

# Chart 2 -- Nighttime percentage by user
axes[0, 1].bar(user_profile["user_id"],
               user_profile["nighttime_pct"],
               color=["red" if p > 22 else "steelblue"
                      for p in user_profile["nighttime_pct"]])
axes[0, 1].axhline(y=22, color="orange", linestyle="--", label="22% threshold")
axes[0, 1].set_title("Nighttime Query Percentage by User")
axes[0, 1].set_ylabel("Nighttime %")
axes[0, 1].tick_params(axis="x", rotation=45)
axes[0, 1].legend()

# Chart 3 -- Scatter plot: total queries vs CBRN percentage
colors = ["red" if not v else "steelblue" for v in user_profile["verified"]]
axes[1, 0].scatter(user_profile["total_queries"],
                   user_profile["cbrn_pct"],
                   c=colors, s=100)
for _, row in user_profile.iterrows():
    axes[1, 0].annotate(row["user_id"],
                        (row["total_queries"], row["cbrn_pct"]),
                        textcoords="offset points", xytext=(5, 5))
axes[1, 0].set_title("Total Queries vs CBRN % (Red=Unverified)")
axes[1, 0].set_xlabel("Total Queries")
axes[1, 0].set_ylabel("CBRN %")

# Chart 4 -- CBRN category distribution
cbrn_counts = merged["cbrn_category"].value_counts()
axes[1, 1].pie(cbrn_counts.values,
               labels=cbrn_counts.index,
               autopct="%1.f%%",
               colors=["steelblue", "red", "orange", "green", "purple"])
axes[1, 1].set_title("CBRN Category Distribution")

plt.tight_layout()
plt.savefig("investigation_dashboard.png", dpi=150, bbox_inches="tight")
plt.close()
print("Dashboard saved as investigation_dashboard.png")

